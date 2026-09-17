"""
Loads ml_service/data/generated/dataset.json (see generate_dataset.py) into
the database: author/editor/reviewer accounts, manuscripts with real
bylines and affiliations, reviewer publications, and the simulated
invitation/review history. Also rewrites PaperBridge Demo Data.xlsx with
the fresh accounts and manuscripts.

Deliberately refuses to run if any @demo-paperbridge.test account already
exists -- run `python manage.py wipe_demo_data --confirm` first.

Usage:
    python manage.py seed_ml_dataset [--skip-pdfs]
"""
import json
import os
import zlib
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from apps.manuscripts import storage
from apps.manuscripts.demo_pdf import build_pdf_bytes
from apps.manuscripts.models import Decision, Manuscript, ManuscriptAffiliation, ManuscriptAuthor
from apps.matching.models import ReviewerPublication
from apps.reviews.models import Review, ReviewAssignment
from apps.users.models import UserProfile, UserRole

DATASET_PATH = Path(settings.BASE_DIR).parent / 'ml_service' / 'data' / 'generated' / 'dataset.json'
XLSX_PATH = Path(settings.BASE_DIR).parent / 'PaperBridge Demo Data.xlsx'

STATUS_AGE_DAYS = {
    'submitted': (0, 3),
    'under_review': (5, 20),
    'revisions_requested': (15, 40),
    'accepted': (20, 60),
    'rejected': (20, 60),
    'published': (40, 90),
}

RECOMMENDATION_TO_DECISION = {'accept': Decision.Type.ACCEPT, 'minor': Decision.Type.MINOR,
                              'major': Decision.Type.MAJOR, 'reject': Decision.Type.REJECT}


class Command(BaseCommand):
    help = 'Seed the ML training dataset (authors, reviewers, editors, manuscripts, history).'

    def add_arguments(self, parser):
        parser.add_argument('--skip-pdfs', action='store_true', help='Skip generating/uploading manuscript PDFs.')

    def handle(self, *args, **options):
        User = get_user_model()
        password = os.environ.get('SEED_DEMO_AUTHOR_PASSWORD')
        if not password:
            raise CommandError('SEED_DEMO_AUTHOR_PASSWORD is not set (see backend/.env).')

        if User.objects.filter(email__iendswith='@demo-paperbridge.test').exists():
            raise CommandError(
                'Demo accounts already exist. Run `python manage.py wipe_demo_data --confirm` first.'
            )

        if not DATASET_PATH.exists():
            raise CommandError(f'{DATASET_PATH} not found -- run ml_service/generate_dataset.py first.')

        with open(DATASET_PATH, encoding='utf-8') as f:
            data = json.load(f)

        password_hash = make_password(password)
        rng_now = timezone.now()

        with transaction.atomic():
            author_users = self._create_people(User, data['authors'], password_hash, UserProfile.Role.AUTHOR)
            editor_users = self._create_people(User, data['editors'], password_hash, UserProfile.Role.EDITOR)
            reviewer_users = self._create_reviewers(User, data['reviewers'], password_hash)

            manuscripts = self._create_manuscripts(
                data['manuscripts'], author_users, editor_users, reviewer_users, rng_now,
            )
            self._create_history(data['manuscripts'], manuscripts, reviewer_users, editor_users, rng_now)

        if not options['skip_pdfs']:
            self._upload_pdfs(data['manuscripts'], manuscripts, author_users)

        self._write_xlsx(data, author_users, reviewer_users, manuscripts, password)

        self.stdout.write(self.style.SUCCESS(
            f'Seeded {len(author_users)} authors, {len(editor_users)} editors, '
            f'{len(reviewer_users)} reviewers, {len(manuscripts)} manuscripts.'
        ))

    # -- people -------------------------------------------------------

    def _create_people(self, User, people, password_hash, role):
        users = []
        for p in people:
            user = User.objects.create(
                username=p['email'], email=p['email'], password=password_hash,
                first_name=p['given_name'], last_name=p['family_name'],
            )
            profile = user.profile
            profile.institution = p['institution']
            profile.role = role
            profile.status = UserProfile.Status.ACTIVE
            profile.save(update_fields=['institution', 'role', 'status'])
            UserRole.objects.update_or_create(user=user, role=role, defaults={'status': UserRole.Status.ACTIVE})
            users.append(user)
        return users

    def _create_reviewers(self, User, reviewers, password_hash):
        users = []
        for r in reviewers:
            user = User.objects.create(
                username=r['email'], email=r['email'], password=password_hash,
                first_name=r['given_name'], last_name=r['family_name'],
            )
            profile = user.profile
            profile.institution = r['institution']
            profile.role = UserProfile.Role.REVIEWER
            profile.status = UserProfile.Status.ACTIVE
            profile.availability_status = r['availability_status']
            profile.specialty_tags = r['specialty_tags']
            profile.expertise_areas = r['expertise_areas']
            profile.research_areas = r['research_areas']
            profile.save(update_fields=[
                'institution', 'role', 'status', 'availability_status',
                'specialty_tags', 'expertise_areas', 'research_areas',
            ])
            UserRole.objects.update_or_create(
                user=user, role=UserProfile.Role.REVIEWER, defaults={'status': UserRole.Status.ACTIVE},
            )
            ReviewerPublication.objects.bulk_create([
                ReviewerPublication(
                    reviewer=user, title=pub['title'], abstract=pub['abstract'],
                    year=pub['year'], venue=pub['venue'], keywords=pub['keywords'],
                )
                for pub in r['publications']
            ])
            users.append(user)
        return users

    # -- manuscripts ----------------------------------------------------

    def _create_manuscripts(self, manuscript_data, author_users, editor_users, reviewer_users, now):
        manuscripts = []
        for m in manuscript_data:
            owner = author_users[m['owner_index']]
            manuscript = Manuscript.objects.create(
                owner=owner,
                article_type=m['article_type'],
                title=m['title'],
                abstract=m['abstract'],
                category=m['category'],
                sub_category=m['sub_category'],
                keywords=m['keywords'],
                specialty_tags=m['specialty_tags'],
                file_key='',  # filled in by _upload_pdfs, or left blank with --skip-pdfs
                no_funding=True,
                no_competing=True,
                ethics_na=True,
                agreed_original=True,
                agreed_not_under_review=True,
                agreed_all_approve=True,
                agreed_policies=True,
                status=m['status'],
            )
            manuscripts.append(manuscript)

            order = 0
            for idx in m['author_indices']:
                person = author_users[idx]
                self._add_byline(manuscript, order, person, corresponding=(idx == m['corresponding_index']))
                order += 1
            for r_idx in m.get('reviewer_coauthor_indices', []):
                person = reviewer_users[r_idx]
                self._add_byline(manuscript, order, person, corresponding=False)
                order += 1
            if 'editor_coauthor_index' in m:
                person = editor_users[m['editor_coauthor_index']]
                self._add_byline(manuscript, order, person, corresponding=False)
                order += 1

            lo, hi = STATUS_AGE_DAYS[m['status']]
            # crc32, not hash(): str hashes are salted per process, which made
            # submitted_at differ between otherwise identical seed runs.
            age_days = lo + (zlib.crc32(manuscript.title.encode('utf-8')) % max(hi - lo, 1))
            submitted_at = now - timezone.timedelta(days=age_days)
            update_fields = {'submitted_at': submitted_at, 'updated_at': submitted_at + timezone.timedelta(days=1)}
            if m['status'] == 'published':
                update_fields['published_at'] = now - timezone.timedelta(days=max(age_days - 10, 1))
            Manuscript.objects.filter(pk=manuscript.pk).update(**update_fields)

        return manuscripts

    def _add_byline(self, manuscript, order, person, corresponding):
        # `person` here is a User instance (author/reviewer/editor account) --
        # the byline copies its identity fields rather than FKing to it, same
        # as a real submission would (see ManuscriptAuthor's own docstring).
        author = ManuscriptAuthor.objects.create(
            manuscript=manuscript, order=order,
            given_name=person.first_name, family_name=person.last_name,
            email=person.email, corresponding=corresponding,
        )
        profile = getattr(person, 'profile', None)
        institution = getattr(profile, 'institution', '') if profile else ''
        if institution:
            ManuscriptAffiliation.objects.create(
                author=author, institution=institution,
            )

    # -- history ----------------------------------------------------------

    def _create_history(self, manuscript_data, manuscripts, reviewer_users, editor_users, now):
        for m, manuscript in zip(manuscript_data, manuscripts):
            decisions_seen = []
            for entry in m['assignments']:
                reviewer = reviewer_users[entry['reviewer_index']]
                invited_at = now - timezone.timedelta(hours=entry['invited_hours_ago'])
                respond_by = invited_at + timezone.timedelta(days=10)

                status_map = {
                    'declined': ReviewAssignment.Status.DECLINED,
                    'accepted': ReviewAssignment.Status.ACCEPTED,
                    'submitted': ReviewAssignment.Status.SUBMITTED,
                }
                assignment = ReviewAssignment.objects.create(
                    manuscript=manuscript, reviewer=reviewer,
                    invited_by=editor_users[0],
                    status=status_map[entry['status']],
                    respond_by=respond_by,
                    due_days=entry['due_days'],
                )
                update_fields = {'invited_at': invited_at}
                if entry['status'] == 'declined':
                    update_fields['responded_at'] = invited_at + timezone.timedelta(days=2)
                    update_fields['decline_reason'] = entry['decline_reason']
                elif entry['status'] in ('accepted', 'submitted'):
                    responded_at = invited_at + timezone.timedelta(days=2)
                    due_at = responded_at + timezone.timedelta(days=entry['due_days'])
                    update_fields['responded_at'] = responded_at
                    update_fields['due_at'] = due_at
                ReviewAssignment.objects.filter(pk=assignment.pk).update(**update_fields)

                if entry['status'] == 'submitted':
                    review_data = entry['review']
                    due_at = update_fields['due_at']
                    submitted_at = due_at - timezone.timedelta(days=2) if entry.get('on_time') else due_at + timezone.timedelta(days=3)
                    review = Review.objects.create(
                        assignment=assignment,
                        originality=review_data['originality'], technical=review_data['technical'],
                        clarity=review_data['clarity'], relevance=review_data['relevance'],
                        recommendation=review_data['recommendation'],
                        summary=review_data['summary'], strengths=review_data['strengths'],
                        weaknesses=review_data['weaknesses'],
                    )
                    Review.objects.filter(pk=review.pk).update(submitted_at=submitted_at)
                    decisions_seen.append(review_data['recommendation'])

            if m['status'] in ('revisions_requested', 'accepted', 'rejected', 'published'):
                self._create_decision(manuscript, m['status'], decisions_seen, editor_users, now)

    def _create_decision(self, manuscript, status, recommendations, editor_users, now):
        if status == 'revisions_requested':
            decision_type = Decision.Type.MAJOR if recommendations.count('major') >= recommendations.count('minor') else Decision.Type.MINOR
        elif status == 'accepted' or status == 'published':
            decision_type = Decision.Type.ACCEPT
        else:
            decision_type = Decision.Type.REJECT
        decision = Decision.objects.create(
            manuscript=manuscript, type=decision_type,
            letter=f'Thank you for submitting "{manuscript.title}". Based on the reviews received, '
                   f'the editorial decision is: {decision_type}.',
            decided_by=editor_users[1 % len(editor_users)],
        )
        decided_at = now - timezone.timedelta(days=5)
        Decision.objects.filter(pk=decision.pk).update(decided_at=decided_at)

    # -- storage ------------------------------------------------------------

    def _upload_pdfs(self, manuscript_data, manuscripts, author_users):
        def upload_one(args):
            m, manuscript = args
            authors_line = ', '.join(
                f"{a.given_name} {a.family_name}" for a in manuscript.authors.all()[:4]
            )
            pdf_bytes = build_pdf_bytes(manuscript.title, authors_line, manuscript.abstract, [])
            key = storage.build_key(manuscript.owner_id, 'manuscript', f'{manuscript.pk}.pdf')
            storage.upload_file(__import__('io').BytesIO(pdf_bytes), key, content_type='application/pdf')
            Manuscript.objects.filter(pk=manuscript.pk).update(
                file_key=key, file_name=f'{manuscript.pk}.pdf', file_size=len(pdf_bytes),
            )

        with ThreadPoolExecutor(max_workers=8) as pool:
            list(pool.map(upload_one, zip(manuscript_data, manuscripts)))
        self.stdout.write(self.style.SUCCESS(f'Uploaded {len(manuscripts)} manuscript PDFs.'))

    # -- xlsx -----------------------------------------------------------

    def _write_xlsx(self, data, author_users, reviewer_users, manuscripts, password):
        if not XLSX_PATH.exists():
            self.stdout.write(self.style.WARNING(f'{XLSX_PATH} not found -- skipping xlsx rewrite.'))
            return
        wb = openpyxl.load_workbook(XLSX_PATH)

        def reset_sheet(name, headers):
            if name in wb.sheetnames:
                del wb[name]
            ws = wb.create_sheet(name)
            ws.append(headers)
            return ws

        ws = reset_sheet('Authors', ['User ID', 'Given Name', 'Family Name', 'Email', 'Password', 'Institution', 'City', 'Country'])
        for user, person in zip(author_users, data['authors']):
            ws.append([user.id, person['given_name'], person['family_name'], person['email'], password,
                       person['institution'], person['city'], person['country']])

        ws = reset_sheet('Reviewers', ['User ID', 'Given Name', 'Family Name', 'Email', 'Password', 'Institution',
                                       'City', 'Country', 'Availability', 'Specialty Tags'])
        for user, person in zip(reviewer_users, data['reviewers']):
            ws.append([user.id, person['given_name'], person['family_name'], person['email'], password,
                       person['institution'], person['city'], person['country'],
                       person['availability_status'], ', '.join(person['specialty_tags'])])

        ws = reset_sheet('Manuscripts', ['Manuscript ID', 'Title', 'Category', 'Article Type', 'Status',
                                         'Corresponding Author', 'Submitted At', 'Specialty Tags'])
        for m, manuscript in zip(data['manuscripts'], manuscripts):
            corresponding = manuscript.authors.filter(corresponding=True).first()
            ws.append([
                manuscript.pk, manuscript.title, manuscript.category, manuscript.article_type,
                manuscript.status, corresponding.email if corresponding else '',
                manuscript.submitted_at.isoformat() if manuscript.submitted_at else '',
                ', '.join(manuscript.specialty_tags or []),
            ])

        wb.save(XLSX_PATH)
        self.stdout.write(self.style.SUCCESS(f'Rewrote {XLSX_PATH.name}.'))
