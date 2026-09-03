"""
seed_demo_reviewers.py — add demo reviewer accounts, backfill specialty_tags
onto the 20 existing demo manuscripts (seeded before that field existed), and
refresh 'PaperBridge Demo Data.xlsx' with a Reviewers sheet + a Specialty Tags
column on Manuscripts.

Additive and idempotent — safe to rerun. Does not touch/wipe authors or
manuscripts beyond setting specialty_tags on the titles listed below.

Usage:
    python manage.py seed_demo_reviewers
"""
import os

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError

from apps.manuscripts.models import Manuscript
from apps.users.models import UserProfile, UserRole
from apps.users.taxonomy import SPECIALTY_TAG_LABELS

User = get_user_model()

DEMO_PASSWORD = os.environ.get('SEED_DEMO_AUTHOR_PASSWORD')

XLSX_PATH = os.path.join(os.path.dirname(__file__), '..', '..', '..', '..', '..', 'PaperBridge Demo Data.xlsx')

# ── Reviewers ─────────────────────────────────────────────────────────────
# Two (Camille Dubois, Oluwaseun Adeyemi) deliberately carry no specialty_tags
# — free-text expertise_areas only — so the fallback matching path (built for
# pre-feature/ORCID-created accounts with nothing tagged) has something real
# to demonstrate, not just the primary tag-overlap path.

REVIEWERS = [
    dict(given_name='Wei Ling', family_name='Tan', institution='Universiti Malaya', city='Kuala Lumpur', country='Malaysia', degree='PhD', availability_status='available', specialty_tags=['signal-processing', 'robotics']),
    dict(given_name='Kwame', family_name='Boateng', institution='University of Ghana', city='Accra', country='Ghana', degree='PhD', availability_status='available', specialty_tags=['nlp', 'cybersecurity', 'data-mining']),
    dict(given_name='Elena', family_name='Petrova', institution='Moscow Institute of Physics and Technology', city='Moscow', country='Russia', degree='PhD', availability_status='busy', specialty_tags=['materials-science', 'condensed-matter']),
    dict(given_name='Astrid', family_name='Nilsen', institution='University of Bergen', city='Bergen', country='Norway', degree='PhD', availability_status='available', specialty_tags=['climate-science', 'applied-mathematics', 'statistics']),
    dict(given_name='Marcus', family_name='Reyes', institution='University of the Philippines', city='Quezon City', country='Philippines', degree='PhD', availability_status='available', specialty_tags=['machine-learning', 'deep-learning', 'nlp']),
    dict(given_name='Ines', family_name='Duarte', institution='Universidade de Coimbra', city='Coimbra', country='Portugal', degree='PhD', availability_status='on_leave', specialty_tags=['agriculture', 'sustainability']),
    dict(given_name='Chidi', family_name='Eze', institution='University of Ibadan', city='Ibadan', country='Nigeria', degree='PhD', availability_status='available', specialty_tags=['public-health', 'health-informatics']),
    dict(given_name='Mei', family_name='Zhang', institution='National University of Singapore', city='Singapore', country='Singapore', degree='PhD', availability_status='busy', specialty_tags=['computational-linguistics', 'nlp']),
    dict(given_name='Viktor', family_name='Hoffmann', institution='RWTH Aachen University', city='Aachen', country='Germany', degree='PhD', availability_status='available', specialty_tags=['distributed-systems', 'cybersecurity', 'blockchain-business']),
    dict(given_name='Sana', family_name='Malik', institution='Lahore University of Management Sciences', city='Lahore', country='Pakistan', degree='PhD', availability_status='available', specialty_tags=['computer-vision', 'machine-learning']),
    dict(given_name='Ben', family_name='Carmichael', institution='University of Edinburgh', city='Edinburgh', country='United Kingdom', degree='PhD', availability_status='available', specialty_tags=['distributed-systems', 'software-engineering']),
    dict(given_name='Talia', family_name='Grunberg', institution='Hebrew University of Jerusalem', city='Jerusalem', country='Israel', degree='PhD', availability_status='busy', specialty_tags=['climate-science', 'sustainability']),
    dict(given_name='Ryo', family_name='Tanaka', institution='Osaka University', city='Osaka', country='Japan', degree='PhD', availability_status='available', specialty_tags=['machine-learning', 'biomedical-engineering']),
    dict(given_name='Camille', family_name='Dubois', institution='Sciences Po', city='Paris', country='France', degree='PhD', availability_status='available', specialty_tags=[], expertise_areas='Household economics, education policy, remittances'),
    dict(given_name='Oluwaseun', family_name='Adeyemi', institution='University of Lagos', city='Lagos', country='Nigeria', degree='PhD', availability_status='available', specialty_tags=[], expertise_areas='Comparative politics, judicial institutions, post-communist governance'),
    dict(given_name='Hana', family_name='Kobayashi', institution='University of Tokyo', city='Tokyo', country='Japan', degree='PhD', availability_status='available', specialty_tags=['psychology', 'education']),
    dict(given_name='Lukas', family_name='Bergmann', institution='Technical University of Munich', city='Munich', country='Germany', degree='PhD', availability_status='busy', specialty_tags=['computer-vision', 'machine-learning']),
]
for r in REVIEWERS:
    r['email'] = f"{r['given_name'].lower().replace(' ', '.')}.{r['family_name'].lower().replace(' ', '.')}@demo-paperbridge.test"
    r.setdefault('expertise_areas', ', '.join(SPECIALTY_TAG_LABELS.get(s, s) for s in r['specialty_tags']))

# ── Manuscript specialty-tag backfill, by title ──────────────────────────
# Two titles are intentionally absent (Remittance Flows..., Post-Communist
# Institutional Persistence...) — no clean taxonomy fit, left to demonstrate
# the free-text fallback from the manuscript side too.

MANUSCRIPT_TAGS = {
    'Field-Realistic Evaluation of Low-Power Sensor Fusion Algorithms': ['signal-processing', 'robotics'],
    'Revisiting Sensor Fusion Baselines Under Field Conditions': ['signal-processing', 'robotics'],
    'A Winnowing-Based Approach to Cross-Lingual Plagiarism Detection': ['nlp', 'cybersecurity', 'data-mining'],
    'Thermal Stability of Perovskite Solar Cells Under Cyclic Humidity Stress': ['materials-science', 'condensed-matter'],
    'Fjord Bathymetry Estimation from Sparse Sonar Transects Using Gaussian Processes': ['climate-science', 'applied-mathematics', 'statistics'],
    'Attention Bottlenecks in Long-Document Transformers: A Diagnostic Study': ['nlp', 'deep-learning', 'machine-learning'],
    'Heat-Tolerant Wheat Cultivars: A Multi-Site Field Trial Across the Arabian Peninsula': ['agriculture', 'sustainability'],
    'Explaining Reviewer Disagreement in Peer Review Using Topic-Level Confidence': ['data-mining', 'machine-learning'],
    'Post-Apartheid Spatial Segregation and Access to Primary Healthcare in Cape Town': ['public-health', 'health-informatics'],
    'Code-Switching Detection in Low-Resource Malay-English Social Media Text': ['computational-linguistics', 'nlp'],
    'Byzantine Fault Tolerance Overhead in Permissioned Blockchain Consensus at Scale': ['distributed-systems', 'cybersecurity', 'blockchain-business'],
    'Azulejo Motif Classification Using Convolutional Neural Networks for Heritage Digitization': ['computer-vision', 'machine-learning'],
    'Energy-Aware Task Scheduling for Heterogeneous Edge Computing Clusters': ['distributed-systems', 'software-engineering'],
    'Microplastic Accumulation Gradients in Great Barrier Reef Sediment Cores': ['climate-science', 'sustainability'],
    'Transfer Learning from Protein Language Models to Enzyme Thermostability Prediction': ['machine-learning', 'biomedical-engineering'],
    'Informal Settlement Growth Detection from Satellite Imagery Using Change-Point Segmentation': ['computer-vision', 'climate-science'],
    'Consumer Trust Recovery Strategies Following Data Breach Disclosure: A Field Experiment': ['marketing', 'cybersecurity'],
    'Anxiety Symptom Trajectories in First-Generation University Students: A Longitudinal Cohort Study': ['psychology', 'education'],
}


class Command(BaseCommand):
    help = 'Add demo reviewer accounts, backfill specialty_tags on demo manuscripts, refresh the demo data xlsx.'

    def handle(self, *args, **options):
        if not DEMO_PASSWORD:
            raise CommandError(
                'SEED_DEMO_AUTHOR_PASSWORD is not set. Add it to backend/.env (see .env.example).'
            )

        self.stdout.write(f'Creating/updating {len(REVIEWERS)} demo reviewer accounts ...')
        users = self._create_reviewers()
        self.stdout.write(self.style.SUCCESS(f'  {len(users)} reviewer accounts ready (password: {DEMO_PASSWORD}).'))

        self.stdout.write('Backfilling specialty_tags on existing demo manuscripts ...')
        tagged = self._backfill_manuscript_tags()
        self.stdout.write(self.style.SUCCESS(f'  {tagged} manuscript(s) tagged.'))

        self.stdout.write('Refreshing PaperBridge Demo Data.xlsx ...')
        self._write_xlsx(users)
        self.stdout.write(self.style.SUCCESS(f'  Written to {os.path.abspath(XLSX_PATH)}'))

        self.stdout.write(self.style.SUCCESS('Done.'))

    def _create_reviewers(self):
        users = []
        for r in REVIEWERS:
            user, created = User.objects.get_or_create(
                email=r['email'],
                defaults={
                    'username': r['email'],
                    'first_name': r['given_name'],
                    'last_name': r['family_name'],
                },
            )
            user.set_password(DEMO_PASSWORD)
            user.save()

            try:
                profile = user.profile
            except UserProfile.DoesNotExist:
                profile = UserProfile.objects.create(user=user)
            profile.role = UserProfile.Role.REVIEWER
            profile.status = UserProfile.Status.ACTIVE
            profile.institution = r['institution']
            profile.degree = r['degree']
            profile.availability_status = r['availability_status']
            profile.specialty_tags = r['specialty_tags']
            profile.expertise_areas = r['expertise_areas']
            profile.save()

            UserRole.objects.update_or_create(
                user=user, role=UserProfile.Role.REVIEWER,
                defaults={'status': UserRole.Status.ACTIVE},
            )
            users.append(user)
        return users

    def _backfill_manuscript_tags(self):
        count = 0
        for title, tags in MANUSCRIPT_TAGS.items():
            updated = Manuscript.objects.filter(title=title).update(specialty_tags=tags)
            if updated:
                count += updated
            else:
                self.stdout.write(self.style.WARNING(f'  No manuscript found with title: {title[:60]}'))
        return count

    def _write_xlsx(self, users):
        wb = openpyxl.load_workbook(XLSX_PATH)

        # ── Reviewers sheet ──────────────────────────────────────────────
        if 'Reviewers' in wb.sheetnames:
            del wb['Reviewers']
        ws = wb.create_sheet('Reviewers')
        ws.append([
            'User ID', 'Given Name', 'Family Name', 'Email', 'Password',
            'Institution', 'City', 'Country', 'Availability', 'Specialty Tags',
        ])
        by_email = {u.email: u for u in users}
        for r in REVIEWERS:
            user = by_email[r['email']]
            labels = ', '.join(SPECIALTY_TAG_LABELS.get(s, s) for s in r['specialty_tags']) or f"(none — {r['expertise_areas']})"
            ws.append([
                user.id, r['given_name'], r['family_name'], r['email'], DEMO_PASSWORD,
                r['institution'], r['city'], r['country'], r['availability_status'], labels,
            ])

        # ── Manuscripts sheet: refresh from the live DB, add Specialty Tags ──
        if 'Manuscripts' in wb.sheetnames:
            del wb['Manuscripts']
        ms = wb.create_sheet('Manuscripts')
        ms.append([
            'Manuscript ID', 'Title', 'Category', 'Article Type', 'Status',
            'Corresponding Author', 'Co-Authors', 'Submitted At',
            'Plagiarism Status', 'Similarity %', 'Specialty Tags',
        ])
        for m in Manuscript.objects.all().order_by('id').prefetch_related('authors'):
            authors = list(m.authors.all())
            corresponding = next((a for a in authors if a.corresponding), authors[0] if authors else None)
            co_authors = [a for a in authors if a is not corresponding]
            check = getattr(m, 'plagiarism_check', None)
            tags = ', '.join(SPECIALTY_TAG_LABELS.get(s, s) for s in m.specialty_tags)
            ms.append([
                m.id, m.title, m.category, m.article_type, m.status,
                f'{corresponding.given_name} {corresponding.family_name} ({corresponding.email})' if corresponding else None,
                '; '.join(f'{a.given_name} {a.family_name} ({a.email})' for a in co_authors) or None,
                m.submitted_at.strftime('%Y-%m-%d %H:%M') if m.submitted_at else None,
                check.status if check else None,
                check.similarity_score if check else None,
                tags or None,
            ])

        wb.save(XLSX_PATH)
